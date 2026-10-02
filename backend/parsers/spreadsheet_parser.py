import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from backend.parsers.errors import SpreadsheetError
from backend.utils.text import clean_text


@dataclass
class ParsedColumn:
    """One spreadsheet column: a dance name plus the dancers listed under it."""

    column_number: int  # 1-based, matches the column the user sees in Excel
    dance_name: str  # "" if the first-row cell was blank
    dancers: list[str] = field(default_factory=list)  # cleaned, in order, duplicates kept


def parse_spreadsheet(filename: str, content: bytes) -> list[ParsedColumn]:
    """Read a .csv or .xlsx file and return one ParsedColumn per non-empty column."""
    extension = Path(filename).suffix.lower()

    if extension == ".csv":
        grid = _read_csv(content)
    elif extension == ".xlsx":
        grid = _read_xlsx(content)
    else:
        raise SpreadsheetError(
            f"'{filename}' is not a supported file type. Please upload a .csv or .xlsx file."
        )

    return _grid_to_columns(grid)


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")  # also strips Excel's invisible BOM marker
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")  # older Windows/Excel files
    return list(csv.reader(io.StringIO(text, newline="")))


def _read_xlsx(content: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise SpreadsheetError(
            "This file could not be opened as an Excel (.xlsx) file. "
            "Please check that it is not damaged and try again."
        ) from error

    try:
        sheet = workbook.worksheets[0]  # we only use the first sheet
        return [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _grid_to_columns(grid: list[list[str]]) -> list[ParsedColumn]:
    """Turn rows of cells into columns: first row = dance name, the rest = dancers."""
    if not grid:
        return []

    # Rows can have different lengths, so pad them to make a rectangle.
    width = max(len(row) for row in grid)
    padded = [row + [""] * (width - len(row)) for row in grid]
    header_row, dancer_rows = padded[0], padded[1:]

    columns = []
    for index in range(width):
        dance_name = clean_text(header_row[index])
        dancers = [clean_text(row[index]) for row in dancer_rows]
        dancers = [name for name in dancers if name]  # blank cells are ignored

        if not dance_name and not dancers:
            continue  # completely empty column, usually stray formatting

        columns.append(ParsedColumn(index + 1, dance_name, dancers))

    return columns